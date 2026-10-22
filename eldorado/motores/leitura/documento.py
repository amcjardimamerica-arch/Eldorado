"""Documentos DOCX (com cabeçalhos, rodapés e notas), ODT e PPTX (uma página por slide)."""
from __future__ import annotations

import io

from .leitores import Leitura, Pagina, Registro, sha256


class LeitorDocumento:
    nome = "documento (docx2python, odfpy, python-pptx)"
    tipos = ("docx", "odt", "pptx")

    def ler(self, conteudo: bytes, nome: str | None, registro: Registro) -> Leitura:
        r = Leitura(tipo="documento", leitor=self.nome, sha256=sha256(conteudo))
        from .leitores import detectar_tipo
        tipo = detectar_tipo(conteudo, nome)
        if tipo == "docx":
            from docx2python import docx2python
            with docx2python(io.BytesIO(conteudo)) as d:
                texto = d.text                                 # inclui cabeçalhos, rodapés e notas
                for tabela in d.body:                          # tabela → linhas → células → parágrafos
                    if len(tabela) > 1 and all(isinstance(lin, list) for lin in tabela):
                        r.tabelas.append([["\n".join(c) if isinstance(c, list) else str(c) for c in lin] for lin in tabela])
            r.paginas.append(Pagina(1, texto))
        elif tipo == "odt":
            from odf import teletype, text
            from odf.opendocument import load
            doc = load(io.BytesIO(conteudo))
            r.paginas.append(Pagina(1, "\n".join(teletype.extractText(p) for p in doc.getElementsByType(text.P))))
        else:
            from pptx import Presentation
            for i, slide in enumerate(Presentation(io.BytesIO(conteudo)).slides, 1):
                textos = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
                r.paginas.append(Pagina(i, "\n".join(textos)))
        return r
