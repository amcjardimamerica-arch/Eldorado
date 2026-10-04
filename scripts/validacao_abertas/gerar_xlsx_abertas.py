import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from collections import Counter
d=json.load(open('dados/validacao_abertas/consolidado.json'))
def s(x):
    if x is None: return ''
    if isinstance(x,(list,tuple)): return '; '.join(map(s,x))
    if isinstance(x,dict): return '; '.join(f'{k}: {s(v)}' for k,v in x.items())
    return str(x)
wb=Workbook()
H=PatternFill('solid',fgColor='1F3864');HF=Font(bold=True,color='FFFFFF')
def sheet(ws,cols,rows,widths):
    ws.append(cols)
    for c in ws[1]: c.fill=H;c.font=HF;c.alignment=Alignment(wrap_text=True,vertical='center')
    for r in rows: ws.append(r)
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
ws=wb.active;ws.title='Resumo'
ws.append(['Validação das oportunidades abertas sem estrela de ouro — 03/10/2026']);ws['A1'].font=Font(bold=True,size=14)
ws.append(['Total',len(d)]);ws.append([])
for t,k in (('Grupo','grupo'),('Veredito','veredito'),('Selo validado','selo_validado'),('Base de leitura','base_leitura'),('Elegibilidade AMC','elig')):
    ws.append([t,'Qtd']);ws.cell(ws.max_row,1).font=Font(bold=True);ws.cell(ws.max_row,2).font=Font(bold=True)
    for a,b in sorted(Counter(x[k] for x in d).items(),key=lambda z:-z[1]): ws.append([a,b])
    ws.append([])
ws.column_dimensions['A'].width=60
sheet(wb.create_sheet('Prioridade'),['Ordem','Grupo','UF','ID','Título','Órgão','Selo sistema','Veredito','Selo validado','Base de leitura','Elegibilidade','Início','Fim','Nota rede','Página oficial','Parecer','Próximo passo','Riscos'],
 [[x['ordem'],x['grupo'],x['uf'],x['id'],x['titulo'],x['orgao'],x['selo_sistema'],x['veredito'],x['selo_validado'],x['base_leitura'],x['elig'],x['inicio'],x['fim'],x['nota_rede'],x['pagina_oficial'],s(x['parecer']),s(x['proximo_passo']),s(x['riscos'])] for x in d],
 [7,10,6,14,40,28,10,18,10,24,12,11,11,8,40,70,50,40])
rows=[]
for x in d:
    for k,v in (x['itens'] or {}).items():
        v=v if isinstance(v,dict) else {}
        rows.append([x['ordem'],x['id'],x['titulo'],k,v.get('estado'),s(v.get('valor')),s(v.get('fonte')),s(v.get('trecho')),s(v.get('motivo'))])
sheet(wb.create_sheet('12 pontos'),['Ordem','ID','Título','Ponto','Estado','Valor','Fonte','Trecho','Motivo'],rows,[7,14,40,18,14,50,40,50,30])
pend=[x for x in d if x['base_leitura'] in ('sem fonte oficial lida','classificado pelo título (sem leitura)','leitura parcial') or x['veredito']=='nao_confirmada']
sheet(wb.create_sheet('Pendências de releitura'),['Ordem','Grupo','ID','Título','Veredito','Base de leitura','Observações','Próximo passo'],
 [[x['ordem'],x['grupo'],x['id'],x['titulo'],x['veredito'],x['base_leitura'],s(x['observacoes']),s(x['proximo_passo'])] for x in pend],[7,10,14,40,18,26,60,50])
div=[x for x in d if x['selo_proposto_validador']!=x['selo_validado']]
sheet(wb.create_sheet('Divergências de selo'),['Ordem','ID','Título','Selo proposto pelo validador','Selo validado (regra)','Base de leitura'],
 [[x['ordem'],x['id'],x['titulo'],x['selo_proposto_validador'],x['selo_validado'],x['base_leitura']] for x in div],[7,14,40,18,18,26])
wb.save('docs/relatorios/PLANILHA-OPORTUNIDADES-ABERTAS-SEM-OURO-2026-10-03.xlsx')
print(len(d),len(pend),len(div))
