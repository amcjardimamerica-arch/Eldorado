import json,re,datetime
from collections import Counter,defaultdict
from urllib.parse import urlparse
R='/home/claude/Eldorado2/'
C=json.load(open(R+'dados/validacao_abertas/consolidado.json'))
ITENS=["Objeto","Prazo de inscrição","Resultado","Prazo de recurso","Valor","Órgão / financiador","Território","Esfera","Requisitos","Anexos","Destinação","Área de atuação"]
GN={'1-GO':'1 · Goiás','2-BR':'2 · Brasil (nacional)','3-INT':'3 · Internacional','4-OUTROS':'4 · Demais estados'}
VN={'aberta_valida':'Aberta e válida','aberta_fora_abrangencia':'Aberta, fora da abrangência da A.M.C.','encerrada':'Encerrada','nao_e_oportunidade_osc':'Não é oportunidade para OSC','duplicada':'Duplicada','nao_confirmada':'Não confirmada (sem fonte oficial lida)'}
SN={'ouro':'Ouro','prata':'Prata','bronze':'Bronze'}
def cut(t,n): 
    t=re.sub(r'\s+',' ',str(t or '')).strip(); return t if len(t)<=n else t[:n-1]+'…'
def hostof(u): 
    try:return urlparse(str(u).split(' ;')[0]).netloc
    except: return ''
B=[]  # blocos
def h(n,t):B.append(('h',n,t))
def p(t):B.append(('p',t))
def bl(L):B.append(('b',L))
def tb(hd,rows,w=None):B.append(('t',hd,rows,w))
n=len(C); g=Counter(r['grupo'] for r in C)
hoje='03/10/2026'
h(1,'Parecer de validação — oportunidades abertas sem estrela de ouro')
p(f'A.M.C. Jardim América · Projeto Eldorado · leitura de {hoje} · {n} oportunidades, na ordem de prioridade do sistema (Goiás → Brasil → internacional → demais estados; dentro de cada bloco, nota da rede e prazo mais próximo).')
h(2,'1. Resposta curta')
V=Counter(r['veredito'] for r in C)
acion=[r for r in C if r['veredito']=='aberta_valida']
p(f'Das {n} oportunidades abertas que ainda não tinham ouro no mapa, apenas {V["aberta_valida"]} estão abertas e são válidas para a A.M.C. ({sum(1 for r in acion if r["grupo"]=="1-GO")} em Goiás, {sum(1 for r in acion if r["grupo"]=="2-BR")} no Brasil, {sum(1 for r in acion if r["grupo"]=="3-INT")} internacional e {sum(1 for r in acion if r["grupo"]=="4-OUTROS")} em outros estados). Outras {V["aberta_fora_abrangencia"]} estão abertas, mas exigem território, público ou tipo de proponente que a A.M.C. não tem; {V["encerrada"]} já encerraram; {V["nao_e_oportunidade_osc"]} não selecionam OSC (credenciamentos de artistas, empresas, compras); {V["duplicada"]} são duplicatas; e {V["nao_confirmada"]} não puderam ser confirmadas em fonte oficial porque o PNCP e os PDFs do Querido Diário estavam inacessíveis em 03/10/2026. Essas últimas ficam como pendência de releitura, sem nenhum dado inventado.')
urg=[r for r in sorted(acion,key=lambda r:r['fim'] or '9999') if r['fim'] and r['fim']<='2026-10-31' and r['grupo'] in('1-GO','2-BR')]
seenu=set();ul=[]
for r in urg:
    k=r['pagina_oficial']
    if k in seenu: continue
    seenu.add(k);ul.append(f"{cut(r['titulo'],60)} (até {r['fim'][8:]}/{r['fim'][5:7]})")
p('Ação mais urgente, abertas e válidas em Goiás e no Brasil com prazo até 31/10/2026: '+'; '.join(ul)+'. A fila completa, com as parcialmente elegíveis, está na seção 3.')
p('Fora deste mapa, a coleta dos livros internacionais (pacote anterior) traz janelas abertas: Criança Esperança/UNESCO (05/10 a 08/11/2026), Camargo (até 05/10/2026), Institut français (até 08/10/2026) e Fundo Ecos (edital 50º até 11/11/2026, restrito a PI/MA/BA).')
h(2,'2. Quadro geral')
rows=[]
for k in ('1-GO','2-BR','3-INT','4-OUTROS'):
    c=Counter(r['veredito'] for r in C if r['grupo']==k)
    rows.append([GN[k],str(g[k])]+[str(c[v]) for v in VN])
rows.append(['Total',str(n)]+[str(V[v]) for v in VN])
tb(['Bloco','Total']+['Aberta válida','Aberta fora de abrangência','Encerrada','Não é p/ OSC','Duplicada','Não confirmada'],rows)
rows=[]
for k in ('1-GO','2-BR','3-INT','4-OUTROS'):
    c=Counter(r['selo_validado'] for r in C if r['grupo']==k); b=Counter(r['base_leitura'] for r in C if r['grupo']==k)
    rows.append([GN[k],str(c['ouro']),str(c['prata']),str(c['bronze']),str(b['leitura completa dos 12 pontos']),str(b['leitura parcial']),str(b['sem fonte oficial lida']+b['classificado pelo título (sem leitura)'])])
tb(['Bloco','Ouro validado','Prata','Bronze','Leitura completa','Leitura parcial','Sem leitura oficial'],rows)
p('Regra do ouro validado: página oficial lida, prazo de inscrição com data ou fluxo contínuo, nenhum dos 12 pontos sem leitura e no máximo três pontos "não informados na fonte". Quem tinha proposta de ouro do validador mas não cumpriu a regra foi rebaixado para prata. Sem página oficial, o selo é bronze.')
# 3 acionaveis
h(2,'3. Fila de ação para a A.M.C. (por prazo)')
def key(r): 
    f=r['fim'] or '9999'; return (f if re.match(r'\d{4}-',f or '') else '9999')
tier=[r for r in C if r['veredito']=='aberta_valida' or (r['veredito']=='aberta_fora_abrangencia' and r['grupo'] in('1-GO','2-BR','3-INT') and r['elig'] in('sim','parcial'))]
seen=set();T=[]
for r in sorted(tier,key=key):
    k=r['pagina_oficial'] or r['id']
    if k in seen: continue
    seen.add(k);T.append(r)
rows=[[str(r['ordem']),GN[r['grupo']].split(' · ')[1],cut(r['titulo'],70),r['fim'] or '—',VN[r['veredito']],SN[r['selo_validado']],r['elig']] for r in T]
tb(['Ordem','Bloco','Oportunidade','Prazo','Veredito','Selo','Elegib. A.M.C.'],rows)
p('Registros que apontam para o mesmo edital foram reunidos (por exemplo, o Natal no Parque de Goiânia aparece duas vezes no mapa).')
# 4 detalhe
h(2,'4. Parecer individual das oportunidades abertas válidas e das parcialmente elegíveis em Goiás, Brasil e internacional')
seen=set()
for r in sorted(T,key=lambda r:(r['ordem'])):
    h(3,f"#{r['ordem']} · {cut(r['titulo'],110)}")
    p(f"{GN[r['grupo']]} · {r['orgao'] or '—'} · UF {r['uf'] or 'nacional'} · veredito: {VN[r['veredito']]} · selo do sistema: {SN[r['selo_sistema']]} → selo validado: {SN[r['selo_validado']]} · leitura: {r['base_leitura']}")
    p(f"Prazo: {r['inicio'] or '?'} → {r['fim'] or '?'} · página oficial: {r['pagina_oficial'] or 'não localizada'}")
    p('Parecer: '+r['parecer']); 
    if r['elegibilidade_amc']: p('Elegibilidade da A.M.C.: '+r['elegibilidade_amc'])
    if r['proximo_passo']: p('Próximo passo: '+r['proximo_passo'])
    if r['riscos']: p('Riscos: '+r['riscos'])
    L=[]
    for k in ITENS:
        it=r['itens'][k]; st={'confirmado':'confirmado','nao_informado':'não informado na fonte','dispensado':'dispensado','nao_lido':'não lido'}[it['estado']]
        v=it.get('valor') or it.get('motivo') or ''
        L.append(f"{k} — {st}: {cut(v,260)}"+(f" [{hostof(it.get('fonte'))}]" if it.get('fonte') and it['estado']=='confirmado' else ''))
    bl(L)
    if r['observacoes']: p('Observações da leitura: '+cut(r['observacoes'],600))
# 5 por grupo
h(2,'5. Lista completa na ordem de prioridade do sistema')
p('A lista abaixo traz as 748 oportunidades. Os 12 pontos de cada uma estão na planilha anexa (aba "12 pontos") e no arquivo consolidado.json.')
for k in ('1-GO','2-BR','3-INT','4-OUTROS'):
    h(3,GN[k])
    rows=[[str(r['ordem']),cut(r['titulo'],80),cut(r['orgao'],34),r['fim'] or '—',VN[r['veredito']].split(' (')[0],SN[r['selo_validado']],r['base_leitura'].replace('classificado pelo título (sem leitura)','só título').replace('leitura completa dos 12 pontos','completa').replace('sem fonte oficial lida','sem leitura')] for r in C if r['grupo']==k]
    tb(['Ordem','Oportunidade','Órgão','Prazo','Veredito','Selo','Leitura'],rows)
# 6 pendencias
h(2,'6. Pendências de releitura')
pend=[r for r in C if r['base_leitura'] in('sem fonte oficial lida','classificado pelo título (sem leitura)')]
causa=Counter()
for r in pend:
    t=(r['observacoes']+' '+r['pagina_oficial'].__str__()+' '+(r['parecer'] or '')).lower()
    causa['PNCP indisponível' if ('pncp' in t) else 'Querido Diário / diário oficial sem PDF' if ('querido' in t or 'diário' in t or 'diario' in t) else 'outras causas']+=1
p(f'{len(pend)} oportunidades ficaram sem leitura de fonte oficial. Por causa dominante: '+'; '.join(f'{k}: {v}' for k,v in causa.most_common())+'. Elas não receberam dado inventado: os 12 pontos ficaram "não lidos". A releitura deve ser feita com o PNCP no ar e com o navegador do computador do titular, que abre os PDFs do Querido Diário.')
rows=[[str(r['ordem']),GN[r['grupo']].split(' · ')[1],cut(r['titulo'],80),VN[r['veredito']].split(' (')[0]] for r in pend]
tb(['Ordem','Bloco','Oportunidade','Veredito provisório'],rows)
# 7 divergencias
h(2,'7. Divergências de dados encontradas no mapa')
div=[r for r in C if re.search(r'link_oficial|não corresponde|nao corresponde|outro (município|assunto|edital)|não bate|não batem|duplicat',r['observacoes'].lower())]
p(f'{len(div)} registros têm divergência registrada pelo validador (link oficial de outro assunto ou município, título que não corresponde ao edital, duplicatas, datas que divergem entre notícia e edital). Recomenda-se corrigir o catálogo antes da próxima rodada. As 25 primeiras:')
bl([f"#{r['ordem']} {cut(r['titulo'],70)} — {cut(r['observacoes'],260)}" for r in div[:25]])
# 8 conselho
h(2,'8. Conselho de 7 lentes')
bl(['Extremamente pessimista: 229 oportunidades sem confirmação e 82 classificadas só pelo título mostram que o mapa mistura sinal e ruído; o risco é a A.M.C. perder prazo real por confiar em registro errado (link de outro município, data de vigência tomada como prazo).',
'Pessimista: em "demais estados" quase tudo é credenciamento de artista ou instituição de longa permanência, sem aderência para uma OSC de Goiânia. Gastar leitura ali tem retorno baixo.',
'Levemente pessimista: dos 35 selos ouro validados, só alguns são realmente utilizáveis pela A.M.C.; ouro aqui significa "informação completa", não "vale a pena".',
'Neutro: separar completude do dado de aderência. A fila de ação (seção 3) usa as duas: aberta, válida e elegível. Parâmetros de qualidade: confirmar sempre no edital antes de protocolar; releitura obrigatória das pendências; corrigir links do catálogo; nunca tratar vigência de credenciamento (2027–2036) como prazo.',
'Levemente otimista: em Goiás há oportunidades concretas e próximas (Natal no Parque, R$ 5 milhões, até 26/10; Cidadão Tech 60+, até R$ 1 milhão, até 26/10) com os 12 pontos lidos.',
'Otimista: a validação já separou o que importa em dias; o sistema agora sabe explicar por que cada registro está onde está, com fonte.',
'Extremamente otimista: com PNCP e diários lidos pelo navegador local, a taxa de confirmação sobe e o mapa vira calendário confiável da A.M.C. Decisão do conselho: agir sobre a fila de ação, reler as pendências e corrigir o catálogo.'])
h(2,'9. Método e limites')
bl(['Fontes: só páginas e PDFs oficiais; agregadores e notícias não provam edital. Nada foi estimado: o que não aparece na fonte está "não informado" ou "não lido".',
'Conteúdo lido foi tratado como dado; nenhuma tentativa de injeção de instruções foi detectada.',
'Nenhum formulário foi enviado, nenhum arquivo baixado, nenhum login feito. Páginas com login gov.br ou bloqueio anti-robô não foram contornadas.',
'Limites desta leitura: PNCP fora do ar (503) em 03/10/2026; PDFs do Querido Diário não abrem na aba; vários editais completos em PDF digitalizado.',
'Os 14 itens que já eram ouro no mapa (762 abertas no total) ficaram fora desta validação, por pedido.'])
h(2,'10. Perguntas ao titular para a próxima rodada')
pend=[r for r in C if r['base_leitura'] in('sem fonte oficial lida','classificado pelo título (sem leitura)')]
bl(['A A.M.C. quer disputar o Natal no Parque (R$ 5 milhões, Parque Mutirama)? O objeto é um evento sazonal; a entidade tem experiência comprovável em eventos infantis?',
'O Cidadão Tech 60+ (SECTI-GO) combina com o público atendido hoje? Há atestados de capacidade técnica de inclusão digital?',
'A entidade tem registro vigente no CMDCA/CMAS e certidões em dia para protocolar até 26/10?',
'Quer incluir oportunidades de outros estados só quando houver parceria local (consórcio ou rede)?',
f'Posso priorizar a releitura das {len(pend)} pendências quando o PNCP voltar?'])
# ---- render
md=[]
for b in B:
    if b[0]=='h': md.append('\n'+'#'*b[1]+' '+b[2]+'\n')
    elif b[0]=='p': md.append(b[1]+'\n')
    elif b[0]=='b': md.append('\n'.join('- '+x for x in b[1])+'\n')
    else:
        hd,rows=b[1],b[2]; md.append('| '+' | '.join(hd)+' |\n|'+'---|'*len(hd)+'\n'+'\n'.join('| '+' | '.join(c.replace('|','/') for c in r)+' |' for r in rows)+'\n')
open(R+'docs/relatorios/PARECER-OPORTUNIDADES-ABERTAS-SEM-OURO-2026-10-03.md','w').write('\n'.join(md))
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.section import WD_ORIENT
d=Document(); s=d.sections[0]; s.orientation=WD_ORIENT.LANDSCAPE; s.page_width,s.page_height=s.page_height,s.page_width
for m in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(s,m,Cm(1.6))
d.styles['Normal'].font.name='Calibri'; d.styles['Normal'].font.size=Pt(10)
for b in B:
    if b[0]=='h': d.add_heading(b[2],level=b[1])
    elif b[0]=='p': d.add_paragraph(b[1])
    elif b[0]=='b':
        for x in b[1]: d.add_paragraph(x,style='List Bullet')
    else:
        hd,rows=b[1],b[2]; t=d.add_table(rows=1,cols=len(hd)); t.style='Light Grid Accent 1'
        for i,x in enumerate(hd): t.rows[0].cells[i].text=x
        for r in rows:
            c=t.add_row().cells
            for i,x in enumerate(r): c[i].text=x
        for row in t.rows:
            for c in row.cells:
                for pp in c.paragraphs:
                    for rn in pp.runs: rn.font.size=Pt(8)
        d.add_paragraph()
d.save(R+'docs/relatorios/PARECER-OPORTUNIDADES-ABERTAS-SEM-OURO-2026-10-03.docx')
print(len(T),len(pend),len(div),causa)
