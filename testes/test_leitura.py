"""A leitura (o canivete suíço). Todos os arquivos de teste são gerados na hora: nenhum PDF entra no repositório."""
import io
import json
import zipfile
from datetime import date
from pathlib import Path

import pymupdf
import pytest

from eldorado.motores.leitura import expurgo
from eldorado.motores.leitura import normalizar as N
from eldorado.motores.leitura.leitores import SemLeitor, detectar_tipo, padrao
from eldorado.motores.leitura.ocr import OCR, OCRIndisponivel
from eldorado.motores.leitura.pdf import LeitorPDF

REG = padrao()


def pdf_edital(paginas=4, cabecalho_com_valor=False, imagem_na_ultima=True) -> bytes:
    doc = pymupdf.open()
    for i in range(1, paginas + 1):
        pg = doc.new_page(width=595, height=842)
        pg.insert_text((40, 30), "22/10/2026 10:00 SEI/GO - 123456 - Edital")                       # marca de impressão (cabeçalho)
        if cabecalho_com_valor:
            pg.insert_text((300, 30), "Total do edital: R$ 900.000,00")                              # protegido: nunca sai
        pg.insert_text((260, 825), f"Página {i} de {paginas}")                                       # rodapé
        if imagem_na_ultima and i == paginas:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 60), 0)
            pix.clear_with(200)
            pg.insert_image(pymupdf.Rect(100, 100, 400, 400), pixmap=pix)
            continue
        pg.insert_text((40, 120), f"Seção {i}. As inscrições vão de 01/10/2026 a 30/11/2026.")
        pg.insert_text((40, 140), "Valor por projeto: R$ 50.000,00. CNPJ do órgão: 02.167.849/0001-80.")
        pg.insert_text((40, 160), "Responsável: CPF 123.456.789-09.")
    return doc.tobytes()


def test_pdf_expurgo_seguro_ocr_e_valores_preservados():
    r = REG.ler(pdf_edital(), "edital.pdf")
    assert r.tipo == "pdf" and len(r.paginas) == 4
    assert r.precisa_ocr == [4]                                                       # escaneada: só tinha o carimbo de impressão
    texto = r.texto
    assert "SEI/GO" not in texto and "Página 1 de 4" not in texto                     # cabeçalho e rodapé repetidos saíram
    assert texto.count("R$ 50.000,00") == 3 and texto.count("02.167.849/0001-80") == 3
    assert "123.456.789-09" not in texto and "***" not in texto and "CPF" not in texto   # CPF removido, sem marca nem menção
    assert any("aguardando OCR" in a for a in r.alertas)


def test_bloco_do_cabecalho_com_valor_nunca_sai():
    r = REG.ler(pdf_edital(cabecalho_com_valor=True, imagem_na_ultima=False))
    assert r.texto.count("R$ 900.000,00") == 4                                        # na dúvida, não apaga


def test_pdf_com_menos_de_3_paginas_nao_sofre_expurgo():
    r = REG.ler(pdf_edital(paginas=2, imagem_na_ultima=False))
    assert "SEI/GO" in r.texto


def test_pdf_protegido_e_corrompido_viram_alerta():
    doc = pymupdf.open(stream=pdf_edital(), filetype="pdf")
    protegido = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="segredo", owner_pw="dono")
    assert any("senha" in a for a in REG.ler(protegido).alertas)
    r = REG.ler(b"%PDF-1.4 isto nao e um pdf de verdade")
    assert r.paginas == [] and any("ilegível" in a or "reserva" in a for a in r.alertas)


def test_tabelas_sob_demanda():
    doc = pymupdf.open()
    pg = doc.new_page()
    pg.insert_text((40, 60), "Cronograma das etapas")
    for i, (etapa, data) in enumerate([("Inscrições", "01/10/2026 a 30/11/2026"), ("Resultado", "15/12/2026")]):
        y = 100 + i * 30
        pg.draw_rect(pymupdf.Rect(40, y, 200, y + 30))
        pg.draw_rect(pymupdf.Rect(200, y, 450, y + 30))
        pg.insert_text((45, y + 20), etapa)
        pg.insert_text((205, y + 20), data)
    b = doc.tobytes()
    assert LeitorPDF().ler(b, None, REG).tabelas == []
    tabelas = LeitorPDF(tabelas="uteis").ler(b, None, REG).tabelas
    assert tabelas and ["Inscrições", "01/10/2026 a 30/11/2026"] in tabelas[0]
    with pytest.raises(ValueError):
        LeitorPDF(tabelas="sim")


def test_pagina_real_do_mpt18():
    h = (Path(__file__).parent / "dados/mpt18.html").read_bytes()
    r = REG.ler(h, "https://www.prt18.mpt.mp.br/servicos/destinacoes")
    assert r.tipo == "html" and len(r.links) >= 100 and all(u.startswith("http") for u in r.links)
    assert r.dados["sistema"]["sistema"] == "Joomla" and "Destinação" in r.dados["titulo"]


def test_pagina_com_dados_embutidos_e_wordpress():
    html = """<!doctype html><html><head><title>Editais</title><meta name="generator" content="WordPress 6.4">
    <script type="application/ld+json">{"@type":"Event","name":"Edital de fomento 2026"}</script>
    <script id="__NEXT_DATA__" type="application/json">{"props":{"editais":[{"titulo":"Edital 1"}]}}</script>
    <script>window.__INITIAL_STATE__ = {editais: [{titulo: 'Edital 2', prazo: '30/11/2026'}]};</script></head>
    <body><a href="/editais/1.pdf">Edital 1</a><a href="#topo">topo</a><table><tr><th>Edital</th><th>Prazo</th></tr><tr><td>1</td><td>30/11</td></tr></table></body></html>"""
    r = REG.ler(html.encode(), "https://cultura.go.gov.br/editais")
    assert r.links == ["https://cultura.go.gov.br/editais/1.pdf"] and r.tabelas == [[["Edital", "Prazo"], ["1", "30/11"]]]
    assert r.dados["jsonld"][0]["name"] == "Edital de fomento 2026"
    assert r.dados["embutidos"]["__NEXT_DATA__"]["props"]["editais"][0]["titulo"] == "Edital 1"
    assert r.dados["embutidos"]["__INITIAL_STATE__"]["editais"][0]["titulo"] == "Edital 2"
    assert r.dados["sistema"] == {"sistema": "WordPress", "api": "https://cultura.go.gov.br/wp-json/wp/v2/posts"}


def _docx():
    import docx
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Secretaria de Cultura"
    d.add_paragraph("O edital tem por objeto a seleção de projetos culturais.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Etapa", "Data", "Inscrições", "30/11/2026"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _xlsx():
    import xlsxwriter
    b = io.BytesIO()
    wb = xlsxwriter.Workbook(b)
    ws = wb.add_worksheet("Créditos")
    for i, linha in enumerate([["Proponente", "Valor"], ["Associação A", 15000], ["Associação B", 2500.5]]):
        ws.write_row(i, 0, linha)
    wb.close()
    return b.getvalue()


def _pptx():
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Edital de fomento"
    s.placeholders[1].text = "Inscrições até 30/11/2026"
    b = io.BytesIO()
    p.save(b)
    return b.getvalue()


def _odt():
    from odf.opendocument import OpenDocumentText
    from odf.text import P
    d = OpenDocumentText()
    d.text.addElement(P(text="Chamamento público para organizações da sociedade civil."))
    b = io.BytesIO()
    d.write(b)
    return b.getvalue()


def test_documentos_planilhas_e_apresentacoes():
    r = REG.ler(_docx(), "x.docx")
    assert r.tipo == "docx" and "Secretaria de Cultura" in r.texto and "tem por objeto" in r.texto and ["Inscrições", "30/11/2026"] in r.tabelas[0]
    r = REG.ler(_xlsx(), "creditos.xlsx")
    assert r.tipo == "xlsx" and r.tabelas[0][1] == ["Associação A", "15000"] and "[Créditos]" in r.texto
    r = REG.ler(_pptx())
    assert r.tipo == "pptx" and "Inscrições até 30/11/2026" in r.texto
    r = REG.ler(_odt())
    assert r.tipo == "odt" and "sociedade civil" in r.texto


def test_tipo_pelo_conteudo_e_nao_pelo_nome():
    assert detectar_tipo(_docx(), "planilha.xlsx") == "docx"
    assert detectar_tipo(b"<!doctype html><html><body>x</body></html>", "arquivo.pdf") == "html"
    assert detectar_tipo(b'{"a": 1}') == "json" and detectar_tipo(b"<?xml version='1.0'?><rss></rss>") == "xml"
    assert detectar_tipo(b"\x89PNG\r\n\x1a\nresto") == "imagem"


def test_texto_json_csv_e_codificacao():
    r = REG.ler("nome;valor\nAssociação;10\n".encode("latin-1"), "dados.csv")
    assert r.tabelas[0] == [["nome", "valor"], ["Associação", "10"]]
    assert REG.ler(b'{"editais": [1, 2]}').dados["json"] == {"editais": [1, 2]}


def _zip(arquivos):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        for nome, conteudo in arquivos:
            z.writestr(nome, conteudo)
    return b.getvalue()


def test_compactados_e_os_seus_limites():
    r = REG.ler(_zip([("edital.pdf", pdf_edital(paginas=2, imagem_na_ultima=False)), ("anexo.docx", _docx())]))
    assert r.tipo == "zip" and {n for n, _ in r.internos} == {"edital.pdf", "anexo.docx"}
    assert [x.tipo for _, x in r.internos] == ["pdf", "docx"]
    import py7zr
    b = io.BytesIO()
    with py7zr.SevenZipFile(b, "w") as z:
        z.writestr(_docx(), "anexo.docx")
    r = REG.ler(b.getvalue())
    assert r.tipo == "7z" and r.internos[0][1].tipo == "docx"
    muitos = REG.ler(_zip([(f"a{i}.txt", b"texto") for i in range(205)]))
    assert len(muitos.internos) == 200 and any("200 primeiros" in a for a in muitos.alertas)
    aninhado = _zip([("n3.zip", _zip([("n2.zip", _zip([("x.txt", b"fundo")]))]))])
    r = REG.ler(aninhado)
    interno = r.internos[0][1].internos[0][1]
    assert any("níveis" in a for a in interno.alertas) and interno.internos == []


def test_imagem_sem_ocr_e_erro_claro():
    with pytest.raises(SemLeitor):
        REG.ler(b"\x89PNG\r\n\x1a\n" + b"0" * 50)
    try:
        import rapidocr  # noqa: F401
    except ImportError:
        with pytest.raises(OCRIndisponivel):
            OCR()


def test_expurgo_ingenuo_seria_perigoso_e_o_seguro_nao_e():
    """O bloco de valor no meio da página, repetido, não está na faixa do cabeçalho: nunca sai."""
    paginas = [(800, [expurgo.Bloco(10, 30, "Cabeçalho do órgão"), expurgo.Bloco(400, 420, "R$"), expurgo.Bloco(780, 795, f"pág {i}")]) for i in range(5)]
    saida = expurgo.aplicar(paginas)
    assert all("R$" in t and "Cabeçalho" not in t and "pág" not in t for t, _ in saida)


def test_normalizacao():
    assert N.datas("de 01/10/2026 a 30/11/2026; até o dia 4 de outubro; 2026-12-03; 31/02/2026", date(2026, 9, 20)) == [
        date(2026, 10, 1), date(2026, 11, 30), date(2026, 10, 4), date(2026, 12, 3)]
    assert N.datas("até 15 de janeiro", date(2026, 11, 20)) == [date(2027, 1, 15)]           # sem ano, já passou: o ano seguinte
    assert N.datas("22.05.2026 a 30.12.2026; valor 1.500,00") == [date(2026, 5, 22), date(2026, 12, 30)]    # com ponto, sem confundir valores
    assert N.valores("R$ 500.000,00, R$ 100 mil, R$ 1,5 milhão e trinta mil reais") == [500000.0, 100000.0, 1500000.0, 30000.0]
    assert N.extenso("dois milhões e quinhentos mil") == 2500000 and N.extenso("edital") is None
    assert N.cnpjs("02.167.849/0001-80 e 11.111.111/1111-11") == ["02167849000180"]


def test_leitores_registrados_cobrem_os_formatos():
    assert set(REG.tipos()) >= {"pdf", "html", "texto", "json", "xml", "csv", "xlsx", "xls", "ods", "docx", "odt", "pptx", "zip", "7z", "rar"}
    json.dumps(REG.tipos())
